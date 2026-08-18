from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "top_censoring" / "Top_Censoring_Update.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_FILL = "F2F4F7"
MUTED = "666666"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=11, bold=False, color="000000", name="Calibri", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(p, before=0, after=6, line=1.10, alignment=WD_ALIGN_PARAGRAPH.LEFT):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    p.alignment = alignment


def add_body(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p)
    run = p.add_run(text)
    set_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        style_paragraph(p, before=16, after=8)
        run = p.add_run(text)
        set_font(run, size=16, bold=True, color=BLUE)
    else:
        style_paragraph(p, before=10, after=5)
        run = p.add_run(text)
        set_font(run, size=13, bold=True, color=BLUE)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4, line=1.15)
    run = p.add_run(text)
    set_font(run)
    return p


def add_metadata(doc, label, value):
    p = doc.add_paragraph()
    style_paragraph(p, after=2, line=1.0)
    r = p.add_run(f"{label}: ")
    set_font(r, bold=True)
    r = p.add_run(value)
    set_font(r)


def set_table_layout(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_screening_table(doc):
    headers = [
        "Survey",
        "Valid\nrecords",
        "Above LIS\nceiling",
        "Weighted population\nshare",
        "Welfare\nshare",
    ]
    rows = [
        ("Malawi 1997", "10,698", "59", "0.435%", "41.60%"),
        ("Malawi 2004", "52,691", "65", "0.128%", "2.24%"),
        ("Malawi 2019", "50,476", "34", "0.046%", "1.20%"),
        ("Belize 1993", "9,097", "22", "0.246%", "20.09%"),
        ("Belize 1997", "10,122", "11", "0.114%", "18.24%"),
        ("Angola 2000", "10,100", "5", "0.010%", "0.80%"),
        ("Angola 2008", "45,398", "5", "0.004%", "0.15%"),
        ("Ecuador 2006", "77,636", "8", "0.009%", "3.77%"),
        ("Seychelles 1999", "813", "1", "0.152%", "4.34%"),
        ("Zimbabwe 2019", "987,700", "72", "0.013%", "0.54%"),
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_layout(table, [2200, 1450, 1500, 2400, 1810])
    set_repeat_table_header(table.rows[0])
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, after=0, line=1.0)
        run = p.add_run(header)
        set_font(run, size=9, bold=True, color=DARK_BLUE)
    for row_values in rows:
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, row_values)):
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(p, after=0, line=1.0)
            run = p.add_run(value)
            set_font(run, size=9.2)
    return table


def add_key_message(doc):
    table = doc.add_table(rows=1, cols=1)
    set_table_layout(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    style_paragraph(p, after=0, line=1.10)
    lead = p.add_run("Preliminary conclusion. ")
    set_font(lead, bold=True, color=DARK_BLUE)
    rest = p.add_run(
        "For Malawi 1997, a small upper-tail group has a large effect on mean welfare and inequality, while poverty estimates are essentially unchanged."
    )
    set_font(rest)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(header, after=0)
    r = header.add_run("TOP-CENSORING ANALYSIS")
    set_font(r, size=8.5, bold=True, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(footer, after=0)
    r = footer.add_run("Internal working update")
    set_font(r, size=8.5, color=MUTED)

    title = doc.add_paragraph()
    style_paragraph(title, after=4)
    r = title.add_run("Top-Censoring Analysis Update")
    set_font(r, size=23, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=16)
    r = subtitle.add_run("Initial literature review, survey screening, and case-study findings")
    set_font(r, size=13, color=MUTED)

    add_metadata(doc, "To", "[Manager's name]")
    add_metadata(doc, "From", "[Your name]")
    add_metadata(doc, "Date", "August 18, 2026")
    add_metadata(doc, "Subject", "Initial assessment of upper-tail treatments in survey microdata")

    rule = doc.add_paragraph()
    style_paragraph(rule, before=8, after=8)
    p_pr = rule._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    add_heading(doc, "Summary")
    add_body(doc, "I reviewed institutional and model-based approaches to the treatment of extreme upper-tail values and screened the available observed survey microdata to identify surveys suitable for this analysis. The initial evidence indicates that Malawi 1997 and Belize 1993/1997 are the strongest case studies for evaluating potential top-coding rules.")

    add_heading(doc, "1. Literature review")
    add_body(doc, "The review distinguishes four approaches. LIS applies a transparent top-code based on the interquartile range of log income. The CPS applies top-coding primarily for disclosure protection. The EU-SILC literature compares trimming, winsorizing, and robust Pareto-based sensitivity approaches. WID addresses the related but distinct problem of survey underrepresentation of top incomes by combining surveys with tax data and national accounts.")

    add_heading(doc, "2. Survey screening")
    add_body(doc, "I reviewed the surveys available in the shared drive and retained files with individual-level, non-grouped microdata. Grouped or binned files were excluded because they cannot identify individual upper-tail observations. For each eligible survey, welfare was expressed in 2021 PPP USD per person per day and survey weights were used throughout.")
    add_body(doc, "The LIS upper ceiling was calculated as exp[Q3(log y) + 3 x IQR(log y)]. The table reports the number of records above this threshold, their weighted population share, and their share of total welfare.")
    add_screening_table(doc)

    p = doc.add_paragraph()
    style_paragraph(p, before=4, after=8)
    r = p.add_run("Note: ")
    set_font(r, size=9, bold=True, color=MUTED)
    r = p.add_run("The ceiling identifies candidate observations for review; it does not by itself establish that an observation is erroneous.")
    set_font(r, size=9, color=MUTED)

    add_heading(doc, "3. Initial case-study findings")
    add_heading(doc, "Malawi 1997", level=2)
    add_body(doc, "The LIS ceiling is 48.57 PPP USD per person per day. Fifty-nine records exceed this ceiling. They represent 0.43% of the weighted population but account for 41.6% of total welfare. Pareto diagnostics indicate that the two largest observations are especially inconsistent with the fitted upper-tail pattern.")
    add_body(doc, "I compared no adjustment, trimming, winsorizing, LIS top-coding, and Pareto-based treatments. All adjustment methods substantially reduce the influence of the upper tail on mean welfare and inequality measures. Poverty estimates remain essentially unchanged.")
    add_key_message(doc)

    add_heading(doc, "Belize 1993 and Belize 1997", level=2)
    add_body(doc, "Both Belize surveys show a similar, though less extreme, pattern. In Belize 1993, 22 observations above the LIS ceiling represent 0.25% of the weighted population and 20.1% of welfare. In Belize 1997, 11 observations represent 0.11% of the weighted population and 18.2% of welfare. Pareto diagnostics suggest an especially irregular upper tail in Belize 1997, where the two highest values deviate strongly from the fitted Pareto pattern.")

    add_heading(doc, "4. Preliminary interpretation and next steps")
    add_body(doc, "The preliminary evidence supports the LIS log-IQR rule as a transparent baseline top-coding method. Pareto diagnostics are useful as a complementary validation and sensitivity tool, particularly for surveys in which a very small group accounts for a large share of total welfare.")
    add_bullet(doc, "Apply the treatment-comparison exercise used for Malawi 1997 to Belize 1993 and Belize 1997.")
    add_bullet(doc, "Assess whether the same operational rule provides stable and comparable results across the selected surveys.")
    add_bullet(doc, "Document cases that require data-construction review before a final top-coding recommendation is made.")

    doc.core_properties.title = "Top-Censoring Analysis Update"
    doc.core_properties.subject = "Initial literature review, survey screening, and case-study findings"
    doc.core_properties.author = ""
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
